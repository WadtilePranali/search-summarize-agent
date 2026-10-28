import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            return
        
        # Parse table markdown
        rows = []
        for tl in table_lines:
            tl = tl.strip()
            if not tl or tl.startswith('| :---') or tl.startswith('|:---') or tl.startswith('|---'):
                continue
            cells = [c.strip() for c in tl.split('|')[1:-1]]
            if cells:
                rows.append(cells)
        
        if rows:
            num_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            for r_idx, row in enumerate(rows):
                for c_idx, cell_text in enumerate(row):
                    if c_idx < num_cols:
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                        
                        # Style Header
                        if r_idx == 0:
                            set_cell_background(cell, "1F4E79")
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.bold = True
                                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                    r.font.size = Pt(10)
                        else:
                            # Alternating row background
                            bg = "F2F5F8" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg)
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.size = Pt(9.5)
            doc.add_paragraph() # Spacing
        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        code_lines = []
        in_code_block = False

    for line in lines:
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables
        if stripped.startswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table()

        # Handle Headings
        if stripped.startswith("# "):
            p = doc.add_heading(level=1)
            run = p.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("## "):
            p = doc.add_heading(level=2)
            run = p.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            p = doc.add_heading(level=3)
            run = p.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x5B, 0x9B, 0xD5)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            # Check for bold prefixes like - **Text:**
            if bullet_text.startswith("**") and "**" in bullet_text[2:]:
                end_bold = bullet_text.find("**", 2)
                bold_part = bullet_text[2:end_bold]
                rest = bullet_text[end_bold+2:]
                r1 = p.add_run(bold_part)
                r1.font.bold = True
                p.add_run(rest)
            else:
                p.add_run(bullet_text)
            p.paragraph_format.space_after = Pt(2)
        elif stripped.startswith("---"):
            # Horizontal rule spacing
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
        elif stripped:
            # Paragraph
            p = doc.add_paragraph()
            # Simple bold formatting handling
            parts = stripped.split("**")
            for idx, part in enumerate(parts):
                run = p.add_run(part)
                if idx % 2 == 1:
                    run.font.bold = True
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Exported successfully to: {docx_path}")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base_dir, "Experiment_10_Documentation.md")
    docx_file = os.path.join(base_dir, "Experiment_10_Documentation.docx")
    markdown_to_docx(md_file, docx_file)
